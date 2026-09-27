"""把待人工处理的两份清单导出为一个 Excel 工作簿，供四位成员分工填写。

输入（本机 out/audit/，不入库）：
    disagreements.jsonl     222 条两标注器分歧
    agreement_sample.jsonl   50 条一致抽样

输出：
    <out>/BN5212_标签裁决_待填写.xlsx
    四个工作表：说明与分工 · 规则表决 · 分歧裁决 · 一致复核

成员以「成员A/B/C/D」占位，不虚构姓名。填写前请自行替换为实际姓名。

用法：
    python code/audit/export_review_xlsx.py --out out/audit
"""
import argparse
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = 'Arial'
HDR_FILL = PatternFill('solid', fgColor='1F3864')
HDR_FONT = Font(name=FONT, bold=True, color='FFFFFF', size=11)
INPUT_FILL = PatternFill('solid', fgColor='FFFF00')       # 需要填写
TITLE_FONT = Font(name=FONT, bold=True, size=14, color='1F3864')
H2_FONT = Font(name=FONT, bold=True, size=11, color='1F3864')
BODY = Font(name=FONT, size=10)
THIN = Side(style='thin', color='BFBFBF')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MEMBERS = ['成员A', '成员B', '成员C', '成员D']
CHOICES = '"1 阳性,0 明确阴性,-1 不确定,空 未提及"'


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BOX


def write_sheet(ws, headers, widths, rows, label_col, note_col):
    ws.append(headers)
    style_header(ws, 1, len(headers))
    ws.row_dimensions[1].height = 30
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ctx_col = headers.index('报告原文（描述性段落节选）') + 1
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.font = BODY
            cell.border = BOX
            cell.alignment = Alignment(vertical='top', wrap_text=(cell.column == ctx_col))
    for col in (label_col, note_col):
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=col).fill = INPUT_FILL
    dv = DataValidation(type='list', formula1=CHOICES, allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv)
    dv.add('%s2:%s%d' % (get_column_letter(label_col), get_column_letter(label_col), ws.max_row))
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(headers)), ws.max_row)


def assign(n):
    """把 n 条平均分给四个人，返回每条的负责人。"""
    out = []
    for i in range(n):
        out.append(MEMBERS[i * len(MEMBERS) // n])
    return out


def build_intro(ws, n_dis, n_agr):
    ws.column_dimensions['A'].width = 4
    ws.column_dimensions['B'].width = 26
    ws.column_dimensions['C'].width = 96
    lines = [
        ('title', 'BN5212 小组项目 · 肺不张标签人工裁决'),
        ('', ''),
        ('h2', '一、这份表格是干什么的'),
        ('p', '我们的实验要比较「训练时包含 vs 排除某位主治医生病例」两个模型的表现差异。模型训练需要一个标签：每次检查到底有没有肺不张。'),
        ('p', '这个标签不是有人重新看了 X 光片得出的——课程数据集本身不提供官方标签，我们只能从放射科医生写的报告文字里提取。'),
        ('p', '标签是整个项目的地基。如果标签有一成是错的，后面的 AUROC 比较全部跟着歪，而且从结果上完全看不出来。所以必须先量一量标签有多准。'),
        ('', ''),
        ('h2', '二、我们已经做了什么'),
        ('p', '用两个互相独立的标注器各标了一遍全部 3,351 份报告，然后比对：'),
        ('li', 'A：CheXbert —— 斯坦福开源的报告标注模型（Smit et al., 2020），学术界通用。'),
        ('li', 'B：第二标注器 —— 逐条阅读报告原文判定，判定前未查看 A 的任何输出（盲标）。'),
        ('p', '结果：3,351 份里一致 3,129 份，一致率 93.4%；分歧 222 份。'),
        ('p', '分歧不是随机噪声——其中 198 条（89%）只归结为两条规则分歧，见下。'),
        ('', ''),
        ('h2', '三、为什么还需要人来做'),
        ('p', '因为两个标注器都是模型，可能在同一类句式上犯同样的错。那时它们会「一致」，从 93.4% 这个数字里完全看不出来。'),
        ('p', '验证需要一个人类锚点。这就是这份表格存在的理由——也是期中汇报里最能说明我们做了真功夫的部分。'),
        ('p', '注意：这不是让你看片子做诊断（那要放射科医生）。诊断是写报告的那位医生早就做完的，你只需要读懂他写的英文，判断他到底说了什么。'),
        ('', ''),
        ('h2', '四、你需要做的三件事'),
        ('p', '第 1 步：在「规则表决」页对两条规则投票（每人 2 票，约 5 分钟）。'),
        ('p', '第 2 步：在「分歧裁决」页填写分给你的行（每人约 %d 条）。' % (n_dis // 4 + 1)),
        ('p', '第 3 步：在「一致复核」页填写分给你的行（每人约 %d 条）。这一步不可省略，它专门用来抓「两个标注器都错」的情况。' % (n_agr // 4 + 1)),
        ('', ''),
        ('h2', '五、怎么填'),
        ('p', '只填黄色格子。黄色 = 需要你填；其余不要改。'),
        ('p', '标签四选一（单元格有下拉菜单）：'),
        ('li', '1 阳性 —— 报告断言存在肺不张，含 mild / minor / streaky 等程度词'),
        ('li', '-1 不确定 —— 报告明确犹豫：may / could / possible / likely / probably，或并列鉴别「atelectasis or pneumonia」'),
        ('li', '0 明确阴性 —— 报告明确否定，如「no atelectasis」'),
        ('li', '空 未提及 —— 描述段落里根本没提'),
        ('', ''),
        ('h2', '六、填写示例'),
        ('example', ''),
        ('', ''),
        ('h2', '七、两条待表决的规则'),
        ('p', '规则一（涉及 164 条分歧）：「likely atelectasis」这类带对冲语气的断言，算阳性还是不确定？'),
        ('p', '    人工构造示例（非病例原文）：Likely atelectasis.'),
        ('p', '    CheXbert 判阳性，第二标注器判不确定。按现行草案「不确定一律排除」，两种选择会让训练集相差约 164 次检查（约 5%）。'),
        ('p', '规则二（涉及 34 条分歧）：报告里的「volume loss」算不算肺不张？'),
        ('p', '    人工构造示例（非病例原文）：Moderate volume loss.'),
        ('p', '    第二标注器计入阳性，CheXbert 未计入。二者在放射学上高度相关但不是同义词；若纳入，需在报告里写明理由。'),
        ('', ''),
        ('h2', '八、完成后'),
        ('p', '把填好的文件发回，汇总后得出：最终标签规则、两标注器一致率、人工裁决结果、以及「两个标注器都错」的比例。这四项写进期中 PPT 与期末报告的方法部分。'),
        ('p', '在这之前，实验设置（标签规则、选片规则、划分比例、医生纳入条件）无法冻结，训练不能开始。'),
    ]
    r = 2
    for kind, text in lines:
        if kind == 'title':
            ws.cell(row=r, column=2, value=text).font = TITLE_FONT
        elif kind == 'h2':
            ws.cell(row=r, column=2, value=text).font = H2_FONT
        elif kind == 'p':
            c = ws.cell(row=r, column=3, value=text)
            c.font = BODY
            c.alignment = Alignment(wrap_text=True, vertical='top')
            ws.row_dimensions[r].height = max(15, 13 * (len(text) // 45 + 1))
        elif kind == 'li':
            c = ws.cell(row=r, column=3, value='· ' + text)
            c.font = BODY
            c.alignment = Alignment(wrap_text=True, vertical='top')
        elif kind == 'example':
            hdrs = ['编号', 'CheXbert 判定', '第二标注器判定', '报告原文（节选）', '你的判定', '理由/备注']
            vals = ['SYNTHETIC-001', '阳性', '不确定',
                    'Likely atelectasis.',
                    '-1 不确定', '医生用了 likely，属于犹豫，按规则一投不确定']
            for j, (h, v) in enumerate(zip(hdrs, vals)):
                hc = ws.cell(row=r, column=3 + j, value=h)
                hc.font = Font(name=FONT, bold=True, size=9)
                hc.border = BOX
                vc = ws.cell(row=r + 1, column=3 + j, value=v)
                vc.font = Font(name=FONT, size=9, italic=True)
                vc.border = BOX
                vc.alignment = Alignment(wrap_text=True, vertical='top')
                if h in ('你的判定', '理由/备注'):
                    vc.fill = INPUT_FILL
            ws.row_dimensions[r + 1].height = 30
            r += 1
        r += 1
    return r


def build_vote(ws):
    ws.append(['规则', '问题', '选项', MEMBERS[0], MEMBERS[1], MEMBERS[2], MEMBERS[3], '多数意见'])
    style_header(ws, 1, 8)
    ws.append(['规则一', '「likely atelectasis」这类对冲断言算什么？（涉及 164 条分歧）',
               '阳性 / 不确定', '', '', '', '', '=IF(COUNTIF(D2:G2,"阳性")>2,"阳性",IF(COUNTIF(D2:G2,"不确定")>2,"不确定","未达多数"))'])
    ws.append(['规则二', '报告中的「volume loss」算不算肺不张？（涉及 34 条分歧）',
               '算 / 不算', '', '', '', '', '=IF(COUNTIF(D3:G3,"算")>2,"算",IF(COUNTIF(D3:G3,"不算")>2,"不算","未达多数"))'])
    for w, col in zip([10, 58, 16, 12, 12, 12, 12, 14], 'ABCDEFGH'):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, max_row=3, max_col=8):
        for cell in row:
            cell.font = BODY
            cell.border = BOX
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for r in (2, 3):
        for c in range(4, 8):
            ws.cell(row=r, column=c).fill = INPUT_FILL
        ws.row_dimensions[r].height = 34
    dv1 = DataValidation(type='list', formula1='"阳性,不确定"', allow_blank=True, showDropDown=False)
    dv2 = DataValidation(type='list', formula1='"算,不算"', allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv1)
    ws.add_data_validation(dv2)
    dv1.add('D2:G2')
    dv2.add('D3:G3')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--out', default='out/audit')
    a = ap.parse_args()

    dis = [json.loads(l) for l in open(os.path.join(a.out, 'disagreements.jsonl'), encoding='utf-8')]
    agr = [json.loads(l) for l in open(os.path.join(a.out, 'agreement_sample.jsonl'), encoding='utf-8')]

    wb = Workbook()
    intro = wb.active
    intro.title = '说明与分工'
    build_intro(intro, len(dis), len(agr))
    build_vote(wb.create_sheet('规则表决'))

    owners = assign(len(dis))
    rows = [['D-%03d' % (i + 1), owners[i], r['study_id'], r['chexbert'], r['second'],
             r['context'], '', ''] for i, r in enumerate(dis)]
    write_sheet(wb.create_sheet('分歧裁决'),
                ['编号', '负责人', '报告编号', 'CheXbert 判定', '第二标注器判定',
                 '报告原文（描述性段落节选）', '你的判定', '理由/备注'],
                [9, 10, 13, 14, 15, 95, 15, 30], rows, 7, 8)

    owners = assign(len(agr))
    rows = [['S-%03d' % (i + 1), owners[i], r['study_id'], r['both'], r['context'], '', ''
             ] for i, r in enumerate(agr)]
    ws = wb.create_sheet('一致复核')
    write_sheet(ws, ['编号', '负责人', '报告编号', '两标注器一致判定',
                     '报告原文（描述性段落节选）', '你的判定', '是否同意'],
                [9, 10, 13, 16, 95, 15, 14], rows, 6, 7)
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=7,
                value='=IF(F%d="","",IF(ISNUMBER(SEARCH(D%d,F%d)),"一致","不一致"))' % (r, r, r))
        ws.cell(row=r, column=7).fill = PatternFill('solid', fgColor='F2F2F2')

    dest = os.path.join(a.out, 'BN5212_标签裁决_待填写.xlsx')
    wb.save(dest)
    print('已导出:', dest)
    print('  分歧裁决 %d 条 · 一致复核 %d 条 · 规则表决 2 条' % (len(dis), len(agr)))
    return 0


if __name__ == '__main__':
    sys.exit(main())
